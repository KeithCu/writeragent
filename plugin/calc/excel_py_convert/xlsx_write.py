# SPDX-License-Identifier: GPL-3.0-or-later
"""Write Excel OOXML packages for DAG =PY and native Python-in-Excel (_xlws.PY).

Handles atomic ZIP writes, worksheet XML patching, script bank emission,
cell/row ordering in sheetData, and spill range clearing.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import TYPE_CHECKING, Any
from xml.etree import ElementTree as ET  # nosemgrep: use-defused-xml  # local .xlsx ZIP parts
import zipfile

from plugin.calc.excel_py_convert.ooxml_util import (
    CALC_CHAIN_REL_TYPE_SUFFIX,
    CONTENT_TYPES_NS,
    PACKAGE_REL_NS,
    PYTHON_SCRIPT_REL_TYPE,
    SSML_NS,
    find_all,
    find_child,
    local_name,
    parse_and_extract_namespaces,
    rewrite_zip,
    serialize_xml_preserving_namespaces,
    workbook_sheets,
    xml_namespace_lock,
)
from plugin.calc.excel_py_convert.parse_dag_formulas import _shared_strings
from plugin.calc.excel_py_convert.parse_excel_ooxml import (
    _col_row,
    _shared_formula_map,
    _sheet_rels_path,
    shift_xlws_py_formula,
)
from plugin.calc.formula_fill import adjust_a1_formula
from plugin.calc.excel_py_convert.script_bank import (
    CODE_SHEET_NOTE,
    CODE_SHEET_NOTE_CELL,
    CODE_SHEET_PREFIX,
    a1_span_bounds,
    collect_script_bank,
    defined_name_targets_sheet,
    formula_for_converted_cell,
    iter_a1_span,
    report_safety_warnings,
    write_script_bank_openpyxl,
)
from plugin.calc.excel_py_convert.to_excel import (
    assign_script_bank,
    deps_for_xlws_export,
    python_scripts_xml,
    xlws_py_formula,
)

if TYPE_CHECKING:
    from plugin.calc.excel_py_convert.models import ConversionReport, ConvertedCell

log = logging.getLogger(__name__)

_RE_A1 = re.compile(r"^([A-Za-z]+)(\d+)$")


def _xlsx_formula_for_cell(cell: ConvertedCell) -> str:
    """Render a comma-separated OOXML ``=PY`` formula (script bank ref, no Calc sanitizer)."""
    return formula_for_converted_cell(cell, separator=",", excel_escape=True, use_script_bank=True)


def non_anchor_spill_cells(anchor: str, array_ref: str) -> list[str]:
    """Return all coordinates in *array_ref* except the *anchor* cell."""
    if not array_ref:
        return []
    clean_anchor = anchor.replace("$", "").upper()
    return [c for c in iter_a1_span(array_ref) if c.replace("$", "").upper() != clean_anchor]


def _clear_spill_range(ws: Any, anchor: str, array_ref: str) -> None:
    """Clear cached/array result cells in *array_ref*, keeping the anchor for rewrite."""
    bounds = a1_span_bounds(array_ref)
    if bounds is None:
        return
    (col1, row1), (col2, row2) = bounds
    parsed = _a1_row_col(anchor)
    anchor_r, anchor_c = parsed if parsed is not None else (0, 0)
    cells = getattr(ws, "_cells", None)
    if not isinstance(cells, dict):
        raise ValueError("openpyxl worksheet._cells is not a dict")
    for (r, c), cell in list(cells.items()):
        if row1 <= r <= row2 and col1 <= c <= col2 and (r, c) != (anchor_r, anchor_c):
            cell.value = None


def _strip_content_types_python(data: bytes) -> bytes:
    """Remove Python-in-Excel Override elements from [Content_Types].xml."""
    # Parse the XML: Excel and LibreOffice write this file as one line, so line filtering deletes it.
    try:
        root, ns_map = parse_and_extract_namespaces(data)
    except ET.ParseError as exc:
        raise ValueError(f"malformed [Content_Types].xml: {exc}") from exc
    changed = False
    for child in list(root):
        if local_name(child.tag) == "Override":
            part_name = child.attrib.get("PartName", "")
            if (
                part_name in ("/xl/pythonScripts.xml", "/xl/calcChain.xml")
                or part_name.startswith("/xl/pythonScripts/")
            ):
                root.remove(child)
                changed = True
    if not changed:
        return data
    return serialize_xml_preserving_namespaces(root, ns_map)


def _strip_rels_python(data: bytes) -> bytes:
    """Remove Python-in-Excel Relationship elements from .rels files."""
    # Parse the XML: a single-line .rels file loses its whole body under line filtering.
    try:
        root, ns_map = parse_and_extract_namespaces(data)
    except ET.ParseError as exc:
        raise ValueError(f"malformed relationship XML (.rels): {exc}") from exc
    changed = False
    for child in list(root):
        if local_name(child.tag) == "Relationship":
            target = child.attrib.get("Target", "").replace("\\", "/").rstrip("/").split("/")[-1]
            rel_type = child.attrib.get("Type", "")
            if (
                rel_type == PYTHON_SCRIPT_REL_TYPE
                or target == "pythonScripts.xml"
                or rel_type.endswith(CALC_CHAIN_REL_TYPE_SUFFIX)
            ):
                root.remove(child)
                changed = True
    if not changed:
        return data
    return serialize_xml_preserving_namespaces(root, ns_map)


def _strip_python_in_excel_parts(out_path: Path) -> None:
    """Remove obsolete Python-in-Excel package parts after formula rewrite."""
    drop_prefixes = ("xl/pythonScripts/",)
    drop_exact = {"xl/pythonScripts.xml", "xl/calcChain.xml"}

    def _transform(zin: zipfile.ZipFile, zout: zipfile.ZipFile) -> None:
        for info in zin.infolist():
            name = info.filename
            if name in drop_exact or any(name.startswith(p) for p in drop_prefixes):
                continue
            data = zin.read(name)
            if name == "[Content_Types].xml":
                data = _strip_content_types_python(data)
            elif name.endswith(".rels"):
                data = _strip_rels_python(data)
            zout.writestr(info, data)

    rewrite_zip(out_path, out_path, _transform, temp_suffix=".tmpstrip")


def write_dag_formulas_xlsx(source_xlsx: str | Path, report: ConversionReport, out_path: str | Path, *, strip_python_parts: bool = True) -> None:
    """Copy *source_xlsx* and replace successfully converted PY cells with DAG formulas.

    Note: openpyxl may alter or drop unsupported package parts and drawing markup on save.
    Where package fidelity is essential, prefer direct XML patching via write_excel_python_xlsx.

    - Parks rewritten Python on visible ``py_code_<Sheet>`` sheets at the **same A1**
      as each caller when ``len(code) > 1000``; shorter scripts stay inline in ``=PY("…")``.
    - OOXML formulas use **comma** separators (not Calc ``;``).
    - Clears the source array/spill ``ref`` range (except the anchor) so old
      cached results do not block the new spill.
    - Fails closed on unmapped sheet titles (no silent first-sheet fallback).
    - Optionally strips ``xl/pythonScripts.xml`` and related package parts.
    """
    import os
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import IllegalCharacterError

    source_xlsx = Path(source_xlsx)
    out_path = Path(out_path)
    if report.direction != "dag":
        raise ValueError("write_dag_formulas_xlsx requires report.direction == 'dag'")
    wb = load_workbook(source_xlsx)
    tmp_path = out_path.with_name(f"{out_path.name}.tmp_{os.getpid()}")
    try:
        sheet_by_key = {ws.title: ws for ws in wb.worksheets}
        sheet_by_key_lower = {t.lower(): w for t, w in sheet_by_key.items()}
        errors: list[str] = []

        bank, bank_warnings = collect_script_bank(report)
        for w in bank_warnings:
            log.warning("excel_py convert: %s", w)
        for w in report_safety_warnings(report):
            log.warning("excel_py convert: %s", w)
        write_script_bank_openpyxl(wb, bank)

        for cell in report.cells:
            if not cell.converted or not cell.converted_code:
                continue
            ws = sheet_by_key.get(cell.sheet) or sheet_by_key_lower.get(cell.sheet.lower())
            if ws is None:
                errors.append(f"unmapped sheet {cell.sheet!r} for cell {cell.cell}")
                continue
            if cell.array_ref:
                _clear_spill_range(ws, cell.cell, cell.array_ref)
            formula = _xlsx_formula_for_cell(cell)
            try:
                ws[cell.cell] = formula
            except IllegalCharacterError as exc:
                errors.append(f"{cell.sheet}!{cell.cell}: {exc}")

        if errors:
            raise ValueError("write_dag_formulas_xlsx failed:\n" + "\n".join(errors))

        wb.save(tmp_path)
        tmp_path.replace(out_path)
    finally:
        if tmp_path.exists() and tmp_path != out_path:
            try:
                tmp_path.unlink()
            except Exception:
                pass
        # Cleanup: close the openpyxl workbook in try/finally on all errors.
        wb.close()
    if strip_python_parts:
        _strip_python_in_excel_parts(out_path)


def _a1_row_col(a1: str) -> tuple[int, int] | None:
    """Return (row_1_based, col_1_based) for an A1 reference."""
    m = _RE_A1.match((a1 or "").replace("$", "").strip())
    if not m:
        return None
    col_str, row_str = m.group(1).upper(), m.group(2)
    col = 0
    for ch in col_str:
        col = col * 26 + (ord(ch) - 64)
    return int(row_str), col


def _ensure_sheet_data(ws_root: ET.Element) -> ET.Element:
    for child in list(ws_root):
        if local_name(child.tag) == "sheetData":
            return child
    tag = f"{{{SSML_NS}}}sheetData" if ws_root.tag.startswith("{") else "sheetData"
    el = ET.SubElement(ws_root, tag)
    return el


class _WorksheetCellIndex:
    """Index of rows and cells in a worksheet for ascending-order lookup and insertion."""

    def __init__(self, ws_root: ET.Element) -> None:
        self.ws_root: ET.Element = ws_root
        self.sheet_data: ET.Element = _ensure_sheet_data(ws_root)
        self.rows: dict[int, ET.Element] = {}
        self.cells_by_row: dict[ET.Element, dict[int, ET.Element]] = {}
        self._build_index()

    def _build_index(self) -> None:
        current_row_num = 0
        for row in list(self.sheet_data):
            if local_name(row.tag) != "row":
                continue
            r_attr = row.attrib.get("r")
            if r_attr is not None:
                try:
                    current_row_num = int(r_attr)
                except ValueError:
                    current_row_num += 1
            else:
                current_row_num += 1
            self.rows[current_row_num] = row
            row_cells: dict[int, ET.Element] = {}
            current_col_num = 0
            for c in list(row):
                if local_name(c.tag) != "c":
                    continue
                c_r = c.attrib.get("r")
                if c_r:
                    parsed = _a1_row_col(c_r)
                    if parsed is not None:
                        current_col_num = parsed[1]
                    else:
                        current_col_num += 1
                else:
                    current_col_num += 1
                row_cells[current_col_num] = c
            self.cells_by_row[row] = row_cells

    def ensure_cell(self, a1: str) -> ET.Element:
        parsed = _a1_row_col(a1)
        if parsed is None:
            raise ValueError(f"invalid A1: {a1!r}")
        target_row, target_col = parsed

        # 1. Locate or insert <row> in ascending order
        row_el = self.rows.get(target_row)
        if row_el is None:
            row_tag = f"{{{SSML_NS}}}row" if self.sheet_data.tag.startswith("{") else "row"
            row_el = ET.Element(row_tag, attrib={"r": str(target_row)})
            insert_row_idx = len(self.sheet_data)
            current_row_num = 0
            for idx, r in enumerate(list(self.sheet_data)):
                if local_name(r.tag) != "row":
                    continue
                r_attr = r.attrib.get("r")
                if r_attr is not None:
                    try:
                        current_row_num = int(r_attr)
                    except ValueError:
                        current_row_num += 1
                else:
                    current_row_num += 1
                if current_row_num > target_row:
                    insert_row_idx = idx
                    break
            self.sheet_data.insert(insert_row_idx, row_el)
            self.rows[target_row] = row_el
            self.cells_by_row[row_el] = {}

        # 2. Locate or insert <c> in ascending order within row_el
        row_cells = self.cells_by_row[row_el]
        if target_col in row_cells:
            return row_cells[target_col]

        col_str = ""
        n = target_col
        while n > 0:
            n, rem = divmod(n - 1, 26)
            col_str = chr(65 + rem) + col_str
        a1_canonical = f"{col_str}{target_row}"

        c_tag = f"{{{SSML_NS}}}c" if row_el.tag.startswith("{") else "c"
        cell = ET.Element(c_tag, attrib={"r": a1_canonical})

        insert_col_idx = len(row_el)
        current_col_num = 0
        for idx, c in enumerate(list(row_el)):
            if local_name(c.tag) != "c":
                insert_col_idx = idx
                break
            c_r = c.attrib.get("r")
            if c_r:
                parsed_c = _a1_row_col(c_r)
                if parsed_c is not None:
                    current_col_num = parsed_c[1]
                else:
                    current_col_num += 1
            else:
                current_col_num += 1
            if current_col_num > target_col:
                insert_col_idx = idx
                break
        row_el.insert(insert_col_idx, cell)
        row_cells[target_col] = cell
        return cell


def _ensure_cell(ws_root: ET.Element, a1: str, index: _WorksheetCellIndex | None = None) -> ET.Element:
    """Return ``<c r="A1">``, creating row/cell under sheetData in ascending order when missing."""
    if index is not None:
        return index.ensure_cell(a1)
    return _WorksheetCellIndex(ws_root).ensure_cell(a1)


def _set_cell_xlws_formula(
    ws_root: ET.Element,
    a1: str,
    formula: str,
    array_ref: str = "",
    index: _WorksheetCellIndex | None = None,
) -> None:
    """Set formula on cell, clearing stale si, ca, v, is, and t attributes."""
    cell = _ensure_cell(ws_root, a1, index=index)
    # Drop cached value / type so Excel recalculates from the formula.
    for child in list(cell):
        if local_name(child.tag) in ("v", "is"):
            cell.remove(child)
    if "t" in cell.attrib:
        del cell.attrib["t"]
    f = find_child(cell, "f")
    if f is None:
        f_tag = f"{{{SSML_NS}}}f" if cell.tag.startswith("{") else "f"
        f = ET.Element(f_tag)
        insert_idx = len(cell)
        for idx, child in enumerate(list(cell)):
            if local_name(child.tag) in ("v", "is", "extLst"):
                insert_idx = idx
                break
        cell.insert(insert_idx, f)
    body = formula[1:] if formula.startswith("=") else formula
    f.text = body
    if "si" in f.attrib:
        del f.attrib["si"]
    if "ca" in f.attrib:
        del f.attrib["ca"]
    if array_ref:
        f.set("t", "array")
        f.set("ref", array_ref.replace("$", ""))
    else:
        if "t" in f.attrib:
            del f.attrib["t"]
        if "ref" in f.attrib:
            del f.attrib["ref"]


def _clear_spill_xml(ws_root: ET.Element, anchor: str, array_ref: str, index: _WorksheetCellIndex | None = None) -> None:
    """Clear cached/array result cells in *array_ref* from XML, keeping the anchor."""
    if not array_ref:
        return
    bounds = a1_span_bounds(array_ref)
    if bounds is None:
        return
    (col1, row1), (col2, row2) = bounds
    anchor_parsed = _a1_row_col(anchor)
    anchor_r, anchor_c = anchor_parsed if anchor_parsed is not None else (0, 0)

    if index is not None:
        for r in range(row1, row2 + 1):
            row_el = index.rows.get(r)
            if row_el is None:
                continue
            row_cells = index.cells_by_row.get(row_el)
            if not row_cells:
                continue
            for col in range(col1, col2 + 1):
                if (r, col) == (anchor_r, anchor_c):
                    continue
                c = row_cells.get(col)
                if c is not None:
                    for child in list(c):
                        if local_name(child.tag) in ("f", "v", "is"):
                            c.remove(child)
                    if "t" in c.attrib:
                        del c.attrib["t"]
        return

    for c in find_all(ws_root, "c"):
        c_r = c.attrib.get("r")
        if not c_r:
            continue
        parsed = _a1_row_col(c_r)
        if parsed is None:
            continue
        r, col = parsed
        if col1 <= col <= col2 and row1 <= r <= row2:
            if (r, col) != (anchor_r, anchor_c):
                for child in list(c):
                    if local_name(child.tag) in ("f", "v", "is"):
                        c.remove(child)
                if "t" in c.attrib:
                    del c.attrib["t"]


def _patch_content_types(data: bytes, drop_parts: set[str] | None = None) -> bytes:
    """Ensure pythonScripts.xml Override is present and remove Overrides for drop_parts and calcChain."""
    try:
        root, ns_map = parse_and_extract_namespaces(data)
    except ET.ParseError as exc:
        raise ValueError(f"malformed [Content_Types].xml: {exc}") from exc
    norm_drop = {p.lstrip("/") for p in drop_parts} if drop_parts else set()
    norm_drop.add("xl/calcChain.xml")
    for child in list(root):
        if local_name(child.tag) == "Override":
            pn = (child.attrib.get("PartName") or "").lstrip("/")
            if pn in norm_drop:
                root.remove(child)
    has_py = any(
        local_name(c.tag) == "Override" and c.attrib.get("PartName") == "/xl/pythonScripts.xml"
        for c in root
    )
    if not has_py:
        tag = f"{{{CONTENT_TYPES_NS}}}Override" if root.tag.startswith("{") else "Override"
        ET.SubElement(
            root,
            tag,
            attrib={"PartName": "/xl/pythonScripts.xml", "ContentType": "application/xml"},
        )
    return serialize_xml_preserving_namespaces(root, ns_map)


def _is_py_code_bank_sheet(
    zin: zipfile.ZipFile,
    sh_title: str,
    sh_part_name: str,
    defined_names: list[tuple[str, str]],
    shared_strings: list[str],
) -> bool:
    """True if sh_title is a script bank sheet (has defined name or ZZ1 == CODE_SHEET_NOTE)."""
    if any(defined_name_targets_sheet(name, val, sh_title) for name, val in defined_names):
        return True
    try:
        ws_bytes = zin.read(sh_part_name)
        ws_root = ET.fromstring(ws_bytes)
    except Exception:
        return False
    for c in find_all(ws_root, "c"):
        if (c.attrib.get("r") or "").replace("$", "").upper() == CODE_SHEET_NOTE_CELL:
            is_el = find_child(c, "is")
            if is_el is not None:
                txt = "".join(is_el.itertext()).strip()
                if txt == CODE_SHEET_NOTE:
                    return True
            t = (c.attrib.get("t") or "").lower()
            v = find_child(c, "v")
            if v is not None and v.text:
                if t == "s":
                    try:
                        idx = int(v.text)
                        if idx < len(shared_strings) and shared_strings[idx] == CODE_SHEET_NOTE:
                            return True
                    except (ValueError, IndexError):
                        pass
                elif v.text.strip() == CODE_SHEET_NOTE:
                    return True
    return False


def _drop_py_code_sheets_from_workbook(
    wb_xml: bytes,
    rels_xml: bytes,
    drop_rids: set[str],
    drop_titles: set[str] | None = None,
) -> tuple[bytes, bytes]:
    """Remove ``py_code_*`` sheet entries and definedNames from workbook.xml + workbook rels."""
    wb, wb_ns = parse_and_extract_namespaces(wb_xml)
    drop_titles_set = {t.lower() for t in drop_titles} if drop_titles else set()
    for sheets_el in list(wb):
        if local_name(sheets_el.tag) != "sheets":
            continue
        for sh in list(sheets_el):
            if local_name(sh.tag) != "sheet":
                continue
            rid = ""
            for k, v in sh.attrib.items():
                if k.endswith("}id") or k in ("r:id", "id"):
                    rid = v
                    break
            title = sh.attrib.get("name") or ""
            if rid in drop_rids or title.lower() in drop_titles_set:
                sheets_el.remove(sh)

    for dn_container in list(wb):
        if local_name(dn_container.tag) != "definedNames":
            continue
        for dn in list(dn_container):
            if local_name(dn.tag) != "definedName":
                continue
            name_attr = dn.attrib.get("name", "")
            val = "".join(dn.itertext()).strip()
            if any(defined_name_targets_sheet(name_attr, val, t) for t in drop_titles_set):
                dn_container.remove(dn)

    calc_pr = find_child(wb, "calcPr")
    if calc_pr is not None:
        calc_pr.attrib["fullCalcOnLoad"] = "1"
    else:
        tag = f"{{{SSML_NS}}}calcPr" if wb.tag.startswith("{") else "calcPr"
        ET.SubElement(wb, tag, attrib={"fullCalcOnLoad": "1"})

    rels, rels_ns = parse_and_extract_namespaces(rels_xml)

    for rel in list(rels):
        if local_name(rel.tag) != "Relationship":
            continue
        rid = rel.attrib.get("Id") or ""
        rel_type = rel.attrib.get("Type") or ""
        if rid in drop_rids or rel_type.endswith(CALC_CHAIN_REL_TYPE_SUFFIX):
            rels.remove(rel)

    has_py_rel = any(
        local_name(r.tag) == "Relationship"
        and (
            r.attrib.get("Type") == PYTHON_SCRIPT_REL_TYPE
            or r.attrib.get("Target", "").replace("\\", "/").rstrip("/").split("/")[-1] == "pythonScripts.xml"
        )
        for r in rels
    )
    if not has_py_rel:
        rids = [
            int(r.attrib["Id"][3:])
            for r in rels
            if local_name(r.tag) == "Relationship"
            and r.attrib.get("Id", "").startswith("rId")
            and r.attrib["Id"][3:].isdigit()
        ]
        next_rid = f"rId{max(rids, default=0) + 1}"
        tag = f"{{{PACKAGE_REL_NS}}}Relationship" if rels.tag.startswith("{") else "Relationship"
        ET.SubElement(
            rels,
            tag,
            attrib={
                "Id": next_rid,
                "Type": PYTHON_SCRIPT_REL_TYPE,
                "Target": "pythonScripts.xml",
            },
        )

    wb_bytes = serialize_xml_preserving_namespaces(wb, wb_ns)
    rels_bytes = serialize_xml_preserving_namespaces(rels, rels_ns)
    return wb_bytes, rels_bytes


def write_excel_python_xlsx(source_xlsx: str | Path, report: ConversionReport, out_path: str | Path) -> None:
    """Write native Excel Python-in-Excel package (stdlib ZipFile only).

    - Banks ``converted_code`` into ``xl/pythonScripts.xml`` (``xl(%Pn%)`` bodies)
    - Sets each cell to ``_xlfn._xlws.PY(scriptIndex, returnType, deps…)``
    - Patches ``[Content_Types].xml``; strips Calc ``py_code_*`` sheets from the package
    """
    source_xlsx = Path(source_xlsx)
    out_path = Path(out_path)
    if report.direction != "excel":
        raise ValueError("write_excel_python_xlsx requires report.direction == 'excel'")
    scripts, _cells = assign_script_bank(report.cells)
    if not scripts or not any(c.converted and c.script_index >= 0 for c in report.cells):
        raise ValueError("write_excel_python_xlsx: no converted Excel PY cells to write")

    errors: list[str] = []

    with xml_namespace_lock(), zipfile.ZipFile(source_xlsx, "r") as zin:
        names = set(zin.namelist())
        if "xl/workbook.xml" not in names:
            raise ValueError("package missing required part: xl/workbook.xml")
        if "xl/_rels/workbook.xml.rels" not in names:
            raise ValueError("package missing required part: xl/_rels/workbook.xml.rels")

        sheets = workbook_sheets(zin)
        title_to_part = {s.title: s.part_name for s in sheets if s.part_name}
        title_to_part_l = {s.title.lower(): s.part_name for s in sheets if s.part_name}

        by_part: dict[str, list[ConvertedCell]] = {}
        seen_part_cells: set[tuple[str, str]] = set()
        for cell in report.cells:
            if not cell.converted or cell.script_index < 0:
                continue
            part = title_to_part.get(cell.sheet) or title_to_part_l.get(cell.sheet.lower(), "")
            if not part:
                errors.append(f"unmapped sheet {cell.sheet!r}")
                continue
            norm_a1 = cell.cell.replace("$", "").upper()
            if (part, norm_a1) in seen_part_cells:
                raise ValueError(f"duplicate cell {norm_a1} for part {part!r} ({cell.sheet!r})")
            seen_part_cells.add((part, norm_a1))
            by_part.setdefault(part, []).append(cell)

        # 1. Scan worksheet formulas for _xlws.PY not in report's write set
        report_write_set: set[tuple[str, str]] = {
            (c.sheet.lower(), c.cell.replace("$", "").upper())
            for c in report.cells
            if c.converted and c.script_index >= 0
        }
        unwritten_cells: list[str] = []
        for sh in sheets:
            if not sh.part_name or sh.title.startswith(CODE_SHEET_PREFIX):
                continue
            try:
                ws_bytes = zin.read(sh.part_name)
            except KeyError:
                continue
            if b"_xlws.PY" not in ws_bytes and b"_xlws.py" not in ws_bytes:
                continue
            ws_root = ET.fromstring(ws_bytes)
            shared_masters = _shared_formula_map(ws_root)
            for c_el in find_all(ws_root, "c"):
                a1 = c_el.attrib.get("r") or ""
                f_el = find_child(c_el, "f")
                if f_el is None or not a1:
                    continue
                body = "".join(f_el.itertext()).strip()
                if (f_el.attrib.get("t") or "") == "shared" and not body:
                    si = f_el.attrib.get("si") or ""
                    if si in shared_masters:
                        body = shared_masters[si][0]
                if "_xlws.py" in body.lower():
                    coord = a1.replace("$", "").upper()
                    if (sh.title.lower(), coord) not in report_write_set:
                        unwritten_cells.append(f"{sh.title}!{a1}")
        if unwritten_cells:
            raise ValueError(
                "Found _xlws.PY cells not in conversion report write set:\n"
                + "\n".join(sorted(unwritten_cells))
            )

        # 2. Bank sheet discovery via defined names / ZZ1 note
        wb_root = ET.fromstring(zin.read("xl/workbook.xml"))
        defined_names: list[tuple[str, str]] = []
        for dn in find_all(wb_root, "definedName"):
            name = dn.attrib.get("name") or ""
            val = "".join(dn.itertext()).strip()
            defined_names.append((name, val))
        shared_strings = _shared_strings(zin)

        code_sheets = [s for s in sheets if s.title.startswith(CODE_SHEET_PREFIX)]
        drop_parts: set[str] = set()
        drop_rids: set[str] = set()
        drop_titles: set[str] = set()
        for s in code_sheets:
            if not s.part_name:
                continue
            if _is_py_code_bank_sheet(zin, s.title, s.part_name, defined_names, shared_strings):
                drop_parts.add(s.part_name)
                drop_rids.add(s.rel_id)
                drop_titles.add(s.title)

        drop_sheet_rels = {_sheet_rels_path(p) for p in drop_parts}

        # 3. Pre-parse and patch worksheet roots
        patched: dict[str, ET.Element] = {}
        original_namespaces: dict[str, dict[str, str]] = {}
        for part, cells in by_part.items():
            root = patched.get(part)
            if root is None:
                try:
                    root, ns_map = parse_and_extract_namespaces(zin.read(part))
                    original_namespaces[part] = ns_map
                except KeyError:
                    errors.append(f"missing worksheet part {part!r}")
                    continue

            ws_index = _WorksheetCellIndex(root)
            sheet_write_set = {cell.cell.replace("$", "").upper() for cell in cells}
            sheet_title = cells[0].sheet

            # Shared formula unsharing
            shared_groups: dict[str, list[tuple[str, ET.Element, ET.Element]]] = {}
            shared_masters = {}
            for c in find_all(root, "c"):
                a1 = c.attrib.get("r") or ""
                f = find_child(c, "f")
                if f is None or not a1:
                    continue
                if (f.attrib.get("t") or "") == "shared":
                    si = f.attrib.get("si") or ""
                    body = "".join(f.itertext()).strip()
                    if body:
                        shared_masters[si] = (body, a1)
                    shared_groups.setdefault(si, []).append((a1, c, f))

            for si, members in shared_groups.items():
                any_in_write = any(m_a1.replace("$", "").upper() in sheet_write_set for m_a1, _, _ in members)
                if not any_in_write:
                    continue
                all_in_write = all(m_a1.replace("$", "").upper() in sheet_write_set for m_a1, _, _ in members)
                if all_in_write:
                    for _m_a1, _m_c, m_f in members:
                        if "si" in m_f.attrib:
                            del m_f.attrib["si"]
                        if "t" in m_f.attrib:
                            del m_f.attrib["t"]
                        if "ref" in m_f.attrib:
                            del m_f.attrib["ref"]
                else:
                    master_info = shared_masters.get(si)
                    if master_info is None:
                        raise ValueError(f"Shared formula group si={si} in {sheet_title!r} has no master body")
                    master_body, master_a1 = master_info
                    m_row, m_col = _col_row(master_a1)
                    for m_a1, _m_c, m_f in members:
                        if m_a1.replace("$", "").upper() in sheet_write_set:
                            if "si" in m_f.attrib:
                                del m_f.attrib["si"]
                            if "t" in m_f.attrib:
                                del m_f.attrib["t"]
                            if "ref" in m_f.attrib:
                                del m_f.attrib["ref"]
                        else:
                            c_row, c_col = _col_row(m_a1)
                            delta_r = c_row - m_row
                            delta_c = c_col - m_col
                            if "_xlws.py" in master_body.lower():
                                shifted = shift_xlws_py_formula(master_body, delta_r, delta_c)
                                if shifted is None:
                                    raise ValueError(f"Cannot unshare formula for {sheet_title}!{m_a1} in si={si}: shifter refused {master_body!r}")
                                m_f.text = shifted[0].lstrip("=")
                            else:
                                master_clean = master_body.lstrip("=")
                                shifted_text = adjust_a1_formula(master_clean, delta_c, delta_r)
                                if "#REF!" in shifted_text and "#REF!" not in master_clean:
                                    raise ValueError(f"Cannot unshare formula for {sheet_title}!{m_a1} in si={si}: out-of-bounds shift produced #REF!")
                                m_f.text = shifted_text
                            if "si" in m_f.attrib:
                                del m_f.attrib["si"]
                            if "t" in m_f.attrib:
                                del m_f.attrib["t"]
                            if "ref" in m_f.attrib:
                                del m_f.attrib["ref"]

            for cell in cells:
                try:
                    if cell.array_ref:
                        _clear_spill_xml(root, cell.cell, cell.array_ref, index=ws_index)
                    formula = xlws_py_formula(cell.script_index, cell.return_type, deps_for_xlws_export(cell))
                    _set_cell_xlws_formula(root, cell.cell, formula, array_ref=cell.array_ref, index=ws_index)
                except Exception as exc:
                    errors.append(f"{cell.sheet}!{cell.cell}: {exc}")
            patched[part] = root

        if errors:
            raise ValueError("write_excel_python_xlsx failed:\n" + "\n".join(errors))

        scripts_bytes = python_scripts_xml(scripts)
        wb_xml = zin.read("xl/workbook.xml")
        rels_xml = zin.read("xl/_rels/workbook.xml.rels")
        wb_xml, rels_xml = _drop_py_code_sheets_from_workbook(wb_xml, rels_xml, drop_rids, drop_titles=drop_titles)

    def _transform(zin: zipfile.ZipFile, zout: zipfile.ZipFile) -> None:
        written_scripts = False
        for info in zin.infolist():
            name = info.filename
            if name in drop_parts or name == "xl/calcChain.xml":
                continue
            if name in drop_sheet_rels:
                continue
            if name == "xl/pythonScripts.xml":
                zout.writestr(info, scripts_bytes)
                written_scripts = True
                continue
            if name == "[Content_Types].xml":
                data = _patch_content_types(zin.read(name), drop_parts=drop_parts)
                zout.writestr(info, data)
                continue
            if name == "xl/workbook.xml":
                zout.writestr(info, wb_xml)
                continue
            if name == "xl/_rels/workbook.xml.rels":
                zout.writestr(info, rels_xml)
                continue
            if name in patched:
                data = serialize_xml_preserving_namespaces(patched[name], original_namespaces.get(name))
                zout.writestr(info, data)
                continue
            zout.writestr(info, zin.read(name))
        if not written_scripts:
            zout.writestr("xl/pythonScripts.xml", scripts_bytes)

    rewrite_zip(source_xlsx, out_path, _transform, temp_suffix=".tmpexcelpy")
